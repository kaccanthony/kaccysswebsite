"""Export Townhall former staff and rank progression data to SQL seed files."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


SOURCE_AS_OF = "2026-10-09"
RETIRED_DISCORD_IDS = {
    "opboy": "759646190999961621",
}
CURRENT_DISCORD = {
    "yoshi5336": ("1289533889815253073", "yoshi5336."),
    "trickzy": ("775004984352964608", "trickzy_gg"),
    "myself_vijnan": ("753274076323119125", "vijnan_2506"),
    "peerione": ("1355223896261136584", "peerione_"),
    "vlackar": ("479880148842184725", "vlackarzmsk"),
    "zig14zag": ("1172924379135492238", "zig14zag."),
    "tuna": ("786911272351956993", "void_tuna_"),
    "jtrms": ("707604709099307009", "jtrms11"),
    "ykreflex": ("1233819553562103918", "ykkreflex"),
    "romanko": ("928892932625416242", "definitelynotroman"),
    "tchai": ("779274835213156363", "cathaypacific_111"),
    "ericjack": ("1155435768580423700", "ericjackofficial"),
    "martysovec1234": ("775666746093862952", "martys5732"),
    "iamhappyketchup": ("1089179220212715550", "germanketchup"),
    "immediatedivider": ("977611575512944651", "immediatedivider"),
    "rostysman": ("1007312984244224040", "rostysman"),
    "captain_ts08": ("1443891124233048176", "captain_ts08"),
    "marciaafr": ("1226496755764559872", "marciaafr"),
}
RANK_PRIORITY = {
    "Operations Manager": 0,
    "Community Manager": 1,
    "Head Staff": 2,
    "Co-Host Authorized": 3,
    "Event Authorized": 4,
    "Assistant Authorized": 5,
}


def sql_text(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def normalize_rank(value: object) -> str:
    label = str(value or "").replace("‑", "-").lower()
    if label.strip() == "om":
        return "Operations Manager"
    if label.strip() == "cm":
        return "Community Manager"
    if "head staff" in label or "host auth" in label and "co-host" not in label:
        return "Head Staff"
    if "co-host" in label:
        return "Co-Host Authorized"
    if "event" in label:
        return "Event Authorized"
    return "Assistant Authorized"


def parse_period(value: object) -> dict[str, object]:
    text = str(value or "").strip()
    match = re.search(
        r"(?P<start>\d{4}-\d{2}-\d{2})\s+to\s+(?P<end>Present|\d{4}-\d{2}-\d{2})"
        r"\s+\(\s*(?P<days>\d+)\s+days\)",
        text,
        re.IGNORECASE,
    )
    if not match:
        return {
            "started_on": "[TO BE ADDED]",
            "ended_on": "[TO BE ADDED]",
            "duration_days": "[TO BE ADDED]",
        }
    return {
        "started_on": match.group("start"),
        "ended_on": match.group("end"),
        "duration_days": int(match.group("days")),
    }


def read_rows(worksheet, start_row: int, end_row: int) -> list[dict[str, object]]:
    people: list[dict[str, object]] = []
    for row_number in range(start_row, end_row + 1):
        name = worksheet.cell(row_number, 2).value
        if not name:
            continue
        timeline = []
        for column in range(5, 33, 2):
            rank_label = worksheet.cell(row_number, column).value
            period = worksheet.cell(row_number, column + 1).value
            if not rank_label:
                continue
            timeline.append({
                "sequence": len(timeline) + 1,
                "staff_rank": normalize_rank(rank_label),
                **parse_period(period),
            })
        people.append({
            "name": str(name).strip(),
            "days": int(float(worksheet.cell(row_number, 3).value)),
            "timeline": timeline,
        })
    return people


def validate_people(current: list[dict[str, object]], former: list[dict[str, object]]) -> None:
    if len(current) != 18 or len(former) != 19:
        raise ValueError(f"Unexpected Townhall row count: {len(current)} current, {len(former)} former")
    for label, people in (("current", current), ("former", former)):
        names = [str(person["name"]).lower() for person in people]
        if len(names) != len(set(names)):
            raise ValueError(f"Duplicate {label} staff names found")
        for person in people:
            durations = [entry["duration_days"] for entry in person["timeline"]]
            if not durations or not all(isinstance(days, int) for days in durations):
                raise ValueError(f"Incomplete timeline for {person['name']}")
            if sum(durations) != person["days"]:
                raise ValueError(f"Timeline total does not match days_as_staff for {person['name']}")
    missing_current_ids = [
        person["name"] for person in current
        if str(person["name"]).lower() not in CURRENT_DISCORD
    ]
    if missing_current_ids:
        raise ValueError(f"Missing current Discord mapping: {', '.join(missing_current_ids)}")


def write_former_staff(path: Path, people: list[dict[str, object]]) -> None:
    values = []
    placeholder_id = -1001
    for person in people:
        name = str(person["name"])
        staff_id = RETIRED_DISCORD_IDS.get(name.lower())
        if staff_id is None:
            staff_id = str(placeholder_id)
            placeholder_id -= 1
        ranks = [entry["staff_rank"] for entry in person["timeline"]]
        highest_rank = min(ranks, key=RANK_PRIORITY.get) if ranks else "[TO BE ADDED]"
        display_name = "_opboy_" if name.lower() == "opboy" else name
        values.append(
            f"  ({staff_id}, {sql_text(display_name)}, {sql_text(display_name)}, "
            f"'[TO BE ADDED]', NULLIF('[TO BE ADDED]', '[TO BE ADDED]')::bigint, "
            f"{sql_text(highest_rank)}, {person['days']}, 0, false)"
        )
    sql = f"""-- Hall of Fame seed exported from Townhall!B30:AF49 on {SOURCE_AS_OF}.
-- Sheet names are temporary display/Discord names. Replace negative staff_id values
-- and '[TO BE ADDED]' Roblox usernames when the real account details are available.
-- Nationality and flag data are deliberately excluded.

INSERT INTO public.staff_archived (
  staff_id,
  staff_name,
  staff_display_name,
  staff_roblox_name,
  staff_roblox_id,
  staff_rank,
  staff_days,
  staff_num_sesh_attend,
  hide_stats
)
VALUES
{',\n'.join(values)}
ON CONFLICT (staff_id) DO UPDATE SET
  staff_name = EXCLUDED.staff_name,
  staff_display_name = EXCLUDED.staff_display_name,
  staff_roblox_name = CASE
    WHEN public.staff_archived.staff_roblox_name = '[TO BE ADDED]'
      THEN EXCLUDED.staff_roblox_name
    ELSE public.staff_archived.staff_roblox_name
  END,
  staff_roblox_id = COALESCE(EXCLUDED.staff_roblox_id, public.staff_archived.staff_roblox_id),
  staff_rank = EXCLUDED.staff_rank,
  staff_days = EXCLUDED.staff_days,
  hide_stats = EXCLUDED.hide_stats;
"""
    path.write_text(sql, encoding="utf-8", newline="\n")


def write_progression(path: Path, people: list[dict[str, object]]) -> None:
    values = []
    for person in people:
        discord_id, discord_username = CURRENT_DISCORD.get(
            str(person["name"]).lower(),
            ("[TO BE ADDED]", "[TO BE ADDED]"),
        )
        timeline = json.dumps(person["timeline"], ensure_ascii=False, separators=(",", ":"))
        values.append(
            f"  ({sql_text(discord_id)}, {sql_text(discord_username)}, "
            f"{sql_text(str(person['name']))}, {person['days']}, {sql_text(timeline)}::jsonb, DATE '{SOURCE_AS_OF}')"
        )
    sql = f"""-- Current staff rank progression exported from Townhall!B7:AF25 on {SOURCE_AS_OF}.
-- staff_display_name preserves the sheet name. Missing account or timeline values are
-- written as '[TO BE ADDED]' so they can be located and completed later.
-- Nationality and flag data are deliberately excluded.

CREATE TABLE IF NOT EXISTS public.staff_rank_progression (
  discord_id text PRIMARY KEY,
  discord_username text NOT NULL DEFAULT '[TO BE ADDED]',
  staff_display_name text NOT NULL,
  days_as_staff integer NOT NULL,
  role_timeline jsonb NOT NULL DEFAULT '[]'::jsonb,
  source_as_of date NOT NULL,
  CONSTRAINT staff_rank_progression_timeline_array
    CHECK (jsonb_typeof(role_timeline) = 'array')
);

ALTER TABLE public.staff_rank_progression ENABLE ROW LEVEL SECURITY;
REVOKE ALL ON public.staff_rank_progression FROM anon, authenticated;
GRANT SELECT, INSERT, UPDATE, DELETE ON public.staff_rank_progression TO service_role;

INSERT INTO public.staff_rank_progression (
  discord_id,
  discord_username,
  staff_display_name,
  days_as_staff,
  role_timeline,
  source_as_of
)
VALUES
{',\n'.join(values)}
ON CONFLICT (discord_id) DO UPDATE SET
  discord_username = EXCLUDED.discord_username,
  staff_display_name = EXCLUDED.staff_display_name,
  days_as_staff = EXCLUDED.days_as_staff,
  role_timeline = EXCLUDED.role_timeline,
  source_as_of = EXCLUDED.source_as_of;
"""
    path.write_text(sql, encoding="utf-8", newline="\n")


def write_retirement(path: Path) -> None:
    sql = """-- Move the three retired accounts out of staff_roster and into staff_archived.
-- Run after staff_public_data.sql. This migration is transactional and verifies the move.
-- opboy's days and highest rank come from Townhall; the other two use their roster data
-- and are calculated through the source date so their totals do not keep increasing.
-- Nationality and flag data are deliberately excluded.

BEGIN;

WITH retiring(discord_id, override_days, override_rank) AS (
  VALUES
    ('759646190999961621', 307, 'Co-Host Authorized'), -- _opboy_
    ('1378997661562175518', NULL::integer, NULL::text), -- timscrperson
    ('1144620459267211295', NULL::integer, NULL::text)  -- your_average_goose
)
INSERT INTO public.staff_archived (
  staff_id,
  staff_name,
  staff_display_name,
  staff_roblox_name,
  staff_rank,
  staff_days,
  staff_num_sesh_attend,
  hide_stats
)
SELECT
  roster.discord_id::bigint,
  roster.discord_username,
  roster.discord_username,
  '[TO BE ADDED]',
  COALESCE(retiring.override_rank, roster.staff_rank),
  COALESCE(retiring.override_days, GREATEST(DATE '2026-10-09' - roster.staff_joined, 0)),
  COALESCE((public_data.operations ->> 'sessions_attended')::numeric::integer, 0),
  false
FROM public.staff_roster AS roster
JOIN retiring ON retiring.discord_id = roster.discord_id
LEFT JOIN public.staff_public_data AS public_data ON public_data.discord_id = roster.discord_id
ON CONFLICT (staff_id) DO UPDATE SET
  staff_name = EXCLUDED.staff_name,
  staff_display_name = EXCLUDED.staff_display_name,
  staff_rank = EXCLUDED.staff_rank,
  staff_days = EXCLUDED.staff_days,
  staff_num_sesh_attend = EXCLUDED.staff_num_sesh_attend,
  hide_stats = EXCLUDED.hide_stats,
  archived_at = now();

DELETE FROM public.staff_roster
WHERE discord_id IN (
  '759646190999961621',
  '1378997661562175518',
  '1144620459267211295'
);

DELETE FROM public.staff_public_data
WHERE discord_id IN (
  '759646190999961621',
  '1378997661562175518',
  '1144620459267211295'
);

DO $$
BEGIN
  IF EXISTS (
    SELECT 1 FROM public.staff_roster
    WHERE discord_id IN (
      '759646190999961621',
      '1378997661562175518',
      '1144620459267211295'
    )
  ) THEN
    RAISE EXCEPTION 'Retirement verification failed: a retired account remains in staff_roster.';
  END IF;

  IF (
    SELECT count(*) FROM public.staff_archived
    WHERE staff_id IN (
      759646190999961621,
      1378997661562175518,
      1144620459267211295
    )
  ) <> 3 THEN
    RAISE EXCEPTION 'Retirement verification failed: expected all three staff_archived rows.';
  END IF;
END $$;

COMMIT;
"""
    path.write_text(sql, encoding="utf-8", newline="\n")


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: export_staff_archive_data.py WORKBOOK OUTPUT_DIRECTORY")
    workbook_path = Path(sys.argv[1])
    output_directory = Path(sys.argv[2])
    worksheet = load_workbook(workbook_path, data_only=True, read_only=True)["Townhall"]
    current = read_rows(worksheet, 8, 25)
    former = read_rows(worksheet, 31, 49)
    validate_people(current, former)
    output_directory.mkdir(parents=True, exist_ok=True)
    former_path = output_directory / "former_staff_data_seed.sql"
    if former_path.exists() and "(-1001," not in former_path.read_text(encoding="utf-8"):
        print("Preserved manually edited former_staff_data_seed.sql")
    else:
        write_former_staff(former_path, former)
    write_progression(output_directory / "staff_rank_progression_seed.sql", current)
    write_retirement(output_directory / "retire_staff_roster.sql")
    print(f"Exported {len(former)} former staff and {len(current)} current progressions.")


if __name__ == "__main__":
    main()
