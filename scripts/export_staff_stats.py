"""Export cached public metrics from the two department workbooks to SQL.

Usage: python scripts/export_staff_stats.py operations.xlsx community.xlsx output.sql [fallback.json]
Requires openpyxl. It reads only Backend identity/stat columns and Operations
Session logs R:S; it never reads Townhall nationality or flag cells.
"""
import json
import sys
from datetime import date
from pathlib import Path

import openpyxl


def plain(value):
    return str(value).strip() if value is not None else ''


def metric(value):
    return round(float(value), 6) if isinstance(value, (int, float)) else None


def sql_text(value):
    return "'" + value.replace("'", "''") + "'"


def load(path):
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def identities(sheet):
    result = {}
    for row in sheet.iter_rows(min_row=6, max_row=100, min_col=1, max_col=4, values_only=True):
        name, _, discord_id, roblox_id = row
        discord_id = plain(discord_id)
        if not (discord_id.isdigit() and 15 <= len(discord_id) <= 21):
            continue
        result[plain(name).casefold()] = (discord_id, int(roblox_id) if isinstance(roblox_id, (int, float)) else None)
    return result


def main(operations_path, community_path, output_path, fallback_path=None):
    ops, comm = load(operations_path), load(community_path)
    ids = identities(ops['Backend']) | identities(comm['Backend'])
    records = {discord_id: {'roblox_id': roblox_id, 'operations': {}, 'community': {}}
               for discord_id, roblox_id in ids.values()}

    session_counts = {}
    for row in ops['Session logs'].iter_rows(min_row=2, max_row=100, min_col=18, max_col=19, values_only=True):
        name, count = row
        if plain(name) and isinstance(count, (int, float)):
            session_counts[plain(name).casefold()] = int(count)

    op_keys = ['weighted_contribution', 'consistency', 'leadership_ratio', 'power_score', 'power_rank']
    for row in ops['Backend'].iter_rows(min_row=6, max_row=200, min_col=18, max_col=23, values_only=True):
        name, *values = row
        key = plain(name).casefold()
        if not key:
            break
        if key not in ids:
            continue
        data = {field: metric(value) for field, value in zip(op_keys, values) if metric(value) is not None}
        if key in session_counts:
            data['sessions_attended'] = session_counts[key]
        records[ids[key][0]]['operations'] = data

    comm_keys = ['runtime_hours', 'weighted_hosting', 'average_turnout', 'tag_versatility', 'department_power_ranking', 'power_rank']
    # Community's stat summary begins at S31. Stop at the first blank name so
    # later Backend sections cannot replace these six metrics.
    for row_number, row in enumerate(comm['Backend'].iter_rows(min_row=31, max_row=200, min_col=19, max_col=25, values_only=True), start=31):
        name, *values = row
        key = plain(name).casefold()
        if not key:
            break
        if key not in ids:
            continue
        data = {
            field: metric(value) for field, value in zip(comm_keys, values) if metric(value) is not None
        }
        if len(data) != len(comm_keys):
            raise ValueError(f'Community Backend row {row_number} has missing cached stats; recalculate and save the workbook first.')
        records[ids[key][0]]['community'] = data

    if fallback_path and Path(fallback_path).exists():
        previous = json.loads(Path(fallback_path).read_text(encoding='utf-8'))
        for discord_id, record in previous.items():
            records.setdefault(discord_id, {
                'roblox_id': record.get('roblox_id'),
                'operations': record.get('operations', {}),
                'community': record.get('community', {}),
            })

    lines = [
        '-- Cached workbook results, exported without nationality or flag data.',
        '-- Re-export and apply after workbook results are recalculated and saved.',
        'BEGIN;',
    ]
    for discord_id, record in sorted(records.items()):
        roblox_id = str(record['roblox_id']) if record['roblox_id'] is not None else 'NULL'
        operations = sql_text(json.dumps(record['operations'], separators=(',', ':')))
        community = sql_text(json.dumps(record['community'], separators=(',', ':')))
        lines.append(
            'INSERT INTO public.staff_public_data (discord_id, roblox_id, operations, community, source_as_of) '
            f"VALUES ({sql_text(discord_id)}, {roblox_id}, {operations}::jsonb, {community}::jsonb, {sql_text(date.today().isoformat())}) "
            'ON CONFLICT (discord_id) DO UPDATE SET '
            'roblox_id = COALESCE(EXCLUDED.roblox_id, staff_public_data.roblox_id), '
            'operations = staff_public_data.operations || EXCLUDED.operations, '
            'community = staff_public_data.community || EXCLUDED.community, '
            'source_as_of = CASE '
            'WHEN (staff_public_data.operations || EXCLUDED.operations) IS DISTINCT FROM staff_public_data.operations '
            'OR (staff_public_data.community || EXCLUDED.community) IS DISTINCT FROM staff_public_data.community '
            'THEN EXCLUDED.source_as_of ELSE staff_public_data.source_as_of END;'
        )
    lines.append('COMMIT;')
    Path(output_path).write_text('\n'.join(lines) + '\n', encoding='utf-8')
    if fallback_path:
        fallback = {discord_id: {**record, 'discord_id': discord_id, 'source_as_of': date.today().isoformat()}
                    for discord_id, record in sorted(records.items())}
        Path(fallback_path).write_text(json.dumps(fallback, indent=2) + '\n', encoding='utf-8')
    print(f'Exported {len(records)} identities to {output_path}')


if __name__ == '__main__':
    if len(sys.argv) not in (4, 5):
        raise SystemExit(__doc__)
    main(*sys.argv[1:])
